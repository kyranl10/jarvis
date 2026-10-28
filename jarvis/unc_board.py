"""Data internships from the UNC CS + DDSS opportunities board: python -m jarvis.unc_board [OUT_FILE]

The board is a public Google Sheet. Claude's cloud sessions can't reach Google Docs directly, so a
GitHub Actions job (.github/workflows/unc.yml) runs this every couple of hours and commits the
filtered rows to data/unc.json, where the internship scan picks them up. Only the sheet's public
listing fields go in; nothing about Kyran.
"""
import csv
import hashlib
import io
import json
import re
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import requests

# The xlsx export is built live (the /edit page can lag) and, unlike the CSV export, keeps each title's link
URL = "https://docs.google.com/spreadsheets/d/1OjnX9zg77zJ2b8uCGCIcgNtKOfI-ARvyHQo2XDQeJ3s/export?format=xlsx&gid=0"
OUT = Path(__file__).resolve().parent.parent / "data" / "unc.json"
TZ = ZoneInfo("America/New_York")
ADDED_WITHIN = 8  # days: rows added to the sheet this recently are worth a look
MAX_AGE = 7       # days: Kyran's one-week rule, measured from the company's posting date

ROLE = re.compile(r"intern|co-?op", re.I)
DATA = re.compile(r"data|analyt|machine learning|\bML\b|\bAI\b|quant|statist|insight|business intel|business analyst"
                  r"|decision science", re.I)
NOT_DATA = re.compile(r"software engineer|hardware|firmware|embedded|mechanical|electrical|security", re.I)
STILL_DATA = re.compile(r"data scien|analyst|analytics|machine learning", re.I)
US = re.compile(r", ?(A[LKZR]|C[AOT]|D[EC]|FL|GA|HI|I[ADLN]|K[SY]|LA|M[ADEINOST]|N[CDEHJMVY]|O[HKR]|PA|RI|S[CD]|T[NX]"
                r"|UT|V[AT]|W[AIVY])\b|\bRemote\b|United States|\bUSA?\b|Hybrid", re.I)
ABROAD = re.compile(r"\b(UK|Canada|India|China|Japan|Germany|France|Ireland|Singapore|Mexico|London|Toronto)",
                    re.I)
UNDERGRAD = re.compile(r"\b(BS|BA|B\.S\.|Bachelor'?s?|Undergrad\w*)\b", re.I)
GRAD_ONLY = re.compile(r"\b(PhD|Ph\.D\.?|MS|M\.S\.|Master'?s?|MBA|Graduate)\b", re.I)

COLUMNS = {  # our name -> words that identify the sheet's column header
    "title": "position title", "company": "company", "role": "role type", "location": "location",
    "added": "date added", "posted": "date posted", "deadline": "deadline", "program": "applicable program",
}


def _columns(header: list[str]) -> dict[str, int]:
    found = {}
    for name, words in COLUMNS.items():
        for i, h in enumerate(header):
            if words in h.lower() and i not in found.values():
                found[name] = i
                break
    return found


def _full_date(text: str, today: date) -> date | None:
    """'10/14/2026' as is; '9/24' as the most recent such day on or before today (the sheet leaves the year off)."""
    m = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\s*", text or "")
    if not m:
        return None
    month, day = int(m[1]), int(m[2])
    try:
        if m[3]:
            year = int(m[3])
            return date(year + 2000 if year < 100 else year, month, day)
        d = date(today.year, month, day)
        return d if d <= today else date(today.year - 1, month, day)
    except ValueError:
        return None


def _short(d: date) -> str:
    return f"{d:%b} {d.day}"


def _noon_ms(d: date) -> int:
    return int(datetime(d.year, d.month, d.day, 12, tzinfo=timezone.utc).timestamp() * 1000)


def _abroad(location: str) -> bool:
    """'London, UK' or 'Toronto, Canada', but not a bare city like 'Raleigh' (the board sometimes leaves the state off)."""
    return not US.search(location) and ("," in location or bool(ABROAD.search(location)))


def job_id(company: str, title: str) -> str:
    return "unc-" + hashlib.sha1(f"{company}|{title}".lower().encode("utf-8")).hexdigest()[:10]


def grad_only(title: str, program: str = "") -> bool:
    text = f"{title} {program}"
    return bool(GRAD_ONLY.search(text)) and not UNDERGRAD.search(text)


def is_data_role(title: str) -> bool:
    return bool(DATA.search(title)) and (not NOT_DATA.search(title) or bool(STILL_DATA.search(title)))


def parse(text: str, today: date) -> dict:
    """The same as parse_table, from the sheet's CSV export (which has no links)."""
    return parse_table(list(csv.reader(io.StringIO(text))), [], today)


def read_xlsx(content: bytes) -> tuple[list[list[str]], list[str]]:
    """(rows of cell text, each row's first link) from the sheet's xlsx export."""
    from openpyxl import load_workbook  # only the UNC board job installs it
    sheet = load_workbook(io.BytesIO(content), data_only=True).worksheets[0]
    table, links = [], []
    for row in sheet.iter_rows():
        cells = []
        for c in row:
            v = c.value
            cells.append(f"{v.month}/{v.day}/{v.year}" if isinstance(v, (datetime, date)) else "" if v is None else str(v))
        table.append(cells)
        links.append(next((c.hyperlink.target for c in row if c.hyperlink and c.hyperlink.target), ""))
    return table, links


def parse_table(table: list[list[str]], links: list[str], today: date) -> dict:
    """{newest_added, rows}: the board's data internships worth a look, newest posting first."""
    start = next((i for i, r in enumerate(table) if any("company" in c.lower() for c in r)), None)
    if start is None:
        raise ValueError("UNC board header row not found")
    links = links + [""] * (len(table) - len(links))
    col = _columns(table[start])
    need = {"title", "company", "role", "location", "added"}
    if not need <= col.keys():
        raise ValueError(f"UNC board columns changed; missing {sorted(need - col.keys())}")
    cell = lambda r, k: r[col[k]].strip() if k in col and col[k] < len(r) else ""
    rows, skipped, newest, seen = [], [], None, set()
    for r, link in zip(table[start + 1:], links[start + 1:]):
        added = _full_date(cell(r, "added"), today)
        if added and (newest is None or added > newest):
            newest = added
        title, company = cell(r, "title"), cell(r, "company")
        if not (title and company and added) or (today - added).days > ADDED_WITHIN:
            continue
        skip = lambda why: skipped.append(f"{_short(added)} {company} | {title}: {why}")
        if not ROLE.search(cell(r, "role")):
            skip(f"role type {cell(r, 'role')!r}")
            continue
        if not is_data_role(title):
            skip("not a data role")
            continue
        if grad_only(title, cell(r, "program")):
            skip("grad students only")
            continue
        if job_id(company, title) in seen:  # the board sometimes lists a posting twice
            continue
        locations = [l.strip() for l in re.split(r"[;\n]", cell(r, "location")) if l.strip()]
        if locations and all(_abroad(l) for l in locations):
            skip(f"location {cell(r, 'location')!r}")
            continue
        posted = _full_date(cell(r, "posted"), today) or added
        if (today - posted).days > MAX_AGE:
            skip(f"posted {cell(r, 'posted')}, over a week ago")
            continue
        deadline = _full_date(cell(r, "deadline"), today + timedelta(days=330))  # deadlines are ahead, not behind
        link = link or next((c.strip() for c in r if c.strip().startswith("http")), "")
        seen.add(job_id(company, title))
        rows.append({"id": job_id(company, title), "company": company, "title": title, "url": link,
                     "locations": locations[:4], "posted": _noon_ms(posted), "sheetAdded": _short(added),
                     **({"deadline": _short(deadline)} if deadline else {}), "source": "UNC board"})
    rows.sort(key=lambda j: -j["posted"])
    return {"newest_added": _short(newest) if newest else None, "rows": rows, "skipped": skipped}


def main(argv: list[str]) -> None:
    out = Path(argv[0]) if argv else OUT
    resp = requests.get(URL, timeout=60)
    resp.raise_for_status()
    data = parse_table(*read_xlsx(resp.content), datetime.now(TZ).date())
    skipped = data.pop("skipped")
    out.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"UNC board: {len(data['rows'])} rows kept, newest added {data['newest_added']}")
    print("Recent rows left out (public sheet fields only):", *skipped, sep="\n  ")


if __name__ == "__main__":
    main(sys.argv[1:])
